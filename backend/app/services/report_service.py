"""Phase 8 — vendor escalation report + PDF/Excel rendering (spec §4–§7).

The vendor report assembles, for one transaction: machine/model/location,
transaction facts, problem classification, timeline, errors, hardware /
cash / host states, the (validated) AI explanation, the full evidence list
and vendor questions. **Every conclusion references evidence** as
``filename:line (EV-xxx)`` — traceability is structural, not optional.

PDF: reportlab platypus (no system dependencies).
Excel: openpyxl workbook with one sheet per required section.
"""

from __future__ import annotations

import io
import xml.sax.saxutils as sax
from datetime import datetime, timezone

from sqlalchemy.orm import Session

from app.ai.explanation import explanation_service
from app.models.transaction import Transaction
from app.services.diagnostic_service import diagnostic_service
from app.services.hardware_service import hardware_service
from app.services.transaction_service import transaction_service

_REPORT_STAMP = "%Y-%m-%d %H:%M:%S UTC"


def _iso(dt) -> str | None:
    return dt.isoformat() if dt else None


def _fmt_dt(dt) -> str:
    if dt is None:
        return "—"
    if isinstance(dt, str):
        return dt.replace("T", " ")[:19]
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def _ev_ref(ev_map: dict, file_id: str | None, line_number: int | None) -> str | None:
    if not file_id or line_number is None:
        return None
    return ev_map.get((file_id, line_number))


def vendor_report(session: Session, txn: Transaction) -> dict:
    """Assemble the complete vendor report (spec §4) with evidence refs."""
    # The explanation (deterministic or external LLM — validated either way).
    explanation = explanation_service.ensure(session, txn)
    payload = explanation["payload"] or {}

    timeline = transaction_service.timeline(session, txn.id)
    from app.schemas.hardware import HardwareTimelineOut

    hw = HardwareTimelineOut.model_validate(
        hardware_service.hardware_payload(session, txn), from_attributes=True
    ).model_dump(mode="json")
    from app.schemas.diagnostics import DiagnosticReportOut

    report = DiagnosticReportOut.model_validate(
        diagnostic_service.report(session, txn), from_attributes=True
    ).model_dump(mode="json")

    machine = txn.machine
    model = txn.machine_model

    # ---- evidence registry: (file_id, line) -> "EV-xxx" --------------------
    from app.ai.digest import EvidenceRegistry, BOUNDS

    ev = EvidenceRegistry()
    for entry in timeline["entries"]:
        raw = entry.get("raw") or {}
        ev.add(raw.get("file_id"), raw.get("line_number"), raw.get("raw_text"), "timeline")
    for m in hw["cash_movements"]:
        ev.add(m.get("log_file_id"), m.get("line_number"), m.get("raw_text"), "cash")
    for s in hw["sensor_events"]:
        ev.add(s.get("log_file_id"), s.get("line_number"), s.get("raw_text"), "sensor")
    for m in hw["motor_events"]:
        ev.add(m.get("log_file_id"), m.get("line_number"), m.get("raw_text"), "motor")
    for g in hw["gate_events"] + hw["shutter_events"]:
        ev.add(g.get("log_file_id"), g.get("line_number"), g.get("raw_text"), "gate")
    for t in hw["transport_events"]:
        ev.add(t.get("log_file_id"), t.get("line_number"), t.get("raw_text"), "transport")
    for f in report["findings"]:
        for evi in f.get("evidence") or []:
            ev.add(evi.get("file_id"), evi.get("line_number"), evi.get("raw_text"), f"rule:{f['rule_id']}")

    from app.models.log_file import LogFile

    file_ids = {item["file_id"] for item in ev.items}
    filenames: dict[str, str] = {}
    if file_ids:
        for row in session.query(LogFile.id, LogFile.original_filename).filter(LogFile.id.in_(file_ids)).all():
            filenames[row[0]] = row[1]
    ev_map = {(i["file_id"], i["line_number"]): i["id"] for i in ev.items}

    def ref(file_id, line_number) -> str:
        ev_id = _ev_ref(ev_map, file_id, line_number)
        name = filenames.get(file_id or "", (file_id or "?")[:8])
        return f"{name}:{line_number}" + (f" ({ev_id})" if ev_id else "")

    # ---- timeline rows ---------------------------------------------------------
    timeline_rows = []
    for entry in timeline["entries"]:
        raw = entry.get("raw") or {}
        timeline_rows.append(
            {
                "timestamp": _fmt_dt(entry.get("timestamp")),
                "event": entry["event"],
                "stage": entry.get("stage"),
                "device": entry.get("device"),
                "severity": entry.get("severity", "INFO"),
                "source": entry.get("source"),
                "not_confirmed": bool(entry.get("not_confirmed")),
                "evidence_ref": ref(raw.get("file_id"), raw.get("line_number"))
                if not entry.get("not_confirmed")
                else None,
                "raw_excerpt": (raw.get("raw_text") or "")[:200] or None,
            }
        )

    # ---- errors -------------------------------------------------------------------
    errors = []
    for entry in timeline["entries"]:
        if entry["event"] in ("ERROR", "VALIDATION_FAILED"):
            detail = entry.get("detail") or {}
            raw = entry.get("raw") or {}
            errors.append(
                {
                    "timestamp": _fmt_dt(entry.get("timestamp")),
                    "error_code": detail.get("error_code"),
                    "description": detail.get("error_description"),
                    "event": entry["event"],
                    "device": entry.get("device"),
                    "evidence_ref": ref(raw.get("file_id"), raw.get("line_number")),
                    "raw_excerpt": (raw.get("raw_text") or "")[:200] or None,
                }
            )

    # ---- hardware / cash / host state blocks --------------------------------------
    hardware_state = {
        "final_classification": hw["faults"][0]["classification"] if hw["faults"] else "NO_EVIDENCE_OF_JAM",
        "faults": [
            {
                "subject": f"{f.get('subject_kind') or ''} {f.get('subject_name') or ''}".strip() or None,
                "classification": f.get("classification"),
                "statement": f.get("statement"),
            }
            for f in hw["faults"]
        ],
        "sensors": [
            {
                "sensor": s.get("sensor"),
                "transition": f"{s.get('previous_state') or '?'}→{s.get('new_state') or '?'}",
                "expected": s.get("expected_state"),
                "actual": s.get("actual_state"),
                "abnormal_duration_ms": s.get("abnormal_duration_ms"),
                "evidence_ref": ref(s.get("log_file_id"), s.get("line_number")),
            }
            for s in hw["sensor_events"]
        ],
        "motors": [
            {
                "motor": m.get("motor"),
                "duration_ms": m.get("duration_ms"),
                "timeout_ms": m.get("timeout_ms"),
                "timed_out": m.get("timed_out"),
                "evidence_ref": ref(m.get("log_file_id"), m.get("line_number")),
            }
            for m in hw["motor_events"]
        ],
        "gates": [
            {
                "kind": g.get("kind"),
                "name": g.get("name"),
                "expected": g.get("expected_state"),
                "actual": g.get("actual_state"),
                "state_mismatch": g.get("state_mismatch"),
                "timed_out": g.get("timed_out"),
                "evidence_ref": ref(g.get("log_file_id"), g.get("line_number")),
            }
            for g in hw["gate_events"] + hw["shutter_events"]
        ],
        "transports": [
            {
                "name": t.get("name"),
                "outcome": t.get("outcome"),
                "timeout_ms": t.get("timeout_ms"),
                "evidence_ref": ref(t.get("log_file_id"), t.get("line_number")),
            }
            for t in hw["transport_events"]
        ],
    }

    cash_state = {
        "final_cash_state": hw["final_cash_state"],
        "movements": [
            {
                "note_id": m.get("note_id"),
                "from_state": m.get("from_state"),
                "to_state": m.get("to_state"),
                "timestamp": _fmt_dt(m.get("timestamp")),
                "device": m.get("device"),
                "confidence": m.get("confidence"),
                "evidence_ref": ref(m.get("log_file_id"), m.get("line_number")),
            }
            for m in hw["cash_movements"]
        ],
    }

    host_state = {
        "request": None,
        "response": None,
        "declined": False,
        "events": [],
    }
    for entry in timeline["entries"]:
        if entry["event"] in ("HOST_REQUEST", "HOST_RESPONSE", "HOST_DECLINED") or (
            entry.get("stage") or ""
        ).startswith("host_"):
            detail = entry.get("detail") or {}
            raw = entry.get("raw") or {}
            host_state["events"].append(
                {
                    "timestamp": _fmt_dt(entry.get("timestamp")),
                    "event": entry["event"],
                    "detail": detail or None,
                    "evidence_ref": ref(raw.get("file_id"), raw.get("line_number"))
                    if not entry.get("not_confirmed")
                    else None,
                }
            )
            if entry["event"] == "HOST_REQUEST":
                host_state["request"] = entry["event"]
            elif entry["event"] == "HOST_RESPONSE":
                host_state["response"] = "RESPONSE_RECEIVED"
            elif entry["event"] == "HOST_DECLINED":
                host_state["response"] = "DECLINED"
                host_state["declined"] = True

    # ---- evidence list -----------------------------------------------------------------
    evidence = [
        {
            "id": item["id"],
            "file": filenames.get(item["file_id"], item["file_id"][:8]),
            "line_number": item["line_number"],
            "raw_excerpt": item["raw_excerpt"],
            "origin": item["origin"],
        }
        for item in ev.items
    ]

    analysis = {
        "classification": report["classification"],
        "diagnosis_class": report["diagnosis_class"],
        "severity": report["severity"],
        "confidence": report["confidence"],
        "engine_summary": report["summary"],
        "findings": [
            {
                "rule_id": f["rule_id"],
                "diagnosis_class": f["diagnosis_class"],
                "severity": f["severity"],
                "confidence": f["confidence"],
                "summary": f["summary"],
                "interpretation": f["interpretation"],
                "possible_causes": f.get("possible_causes") or [],
                "recommended_action": f.get("recommended_action"),
                "evidence_refs": [
                    ref(evi.get("file_id"), evi.get("line_number"))
                    for evi in f.get("evidence") or []
                    if evi.get("file_id") and evi.get("line_number") is not None
                ],
            }
            for f in report["findings"]
        ],
        "ai_explanation": {
            "provider": explanation["provider"],
            "generator": explanation["generator"],
            "model_name": explanation.get("model_name"),
            "digest_sha256": explanation["digest_sha256"],
            "generated_at": explanation["created_at"],
            "technical_summary": payload.get("technical_summary"),
            "root_cause": payload.get("root_cause"),
            "confidence": payload.get("confidence"),
            "possible_causes": payload.get("possible_causes"),
            "caveats": payload.get("caveats"),
            "safety_notes": explanation.get("safety_notes") or [],
        },
    }

    return {
        "generated_at": datetime.now(timezone.utc).strftime(_REPORT_STAMP),
        "report_kind": "VENDOR_ESCALATION",
        "system": "CDM Log Analyzer — Universal GRG CDM Log Analysis",
        "machine": {
            "serial_number": machine.serial_number if machine else None,
            "name": machine.name if machine else None,
            "location": machine.location if machine else None,
            "model": model.code if model else txn.model_code,
            "model_name": model.name if model else None,
            "vendor": model.vendor if model else "GRG Banking",
        },
        "transaction": {
            "id": txn.id,
            "transaction_id": txn.transaction_id,
            "start_time": _iso(txn.start_time),
            "end_time": _iso(txn.end_time),
            "amount": txn.amount,
            "currency": txn.currency,
            "status": txn.status,
            "correlation_confidence": txn.correlation_confidence,
        },
        "problem": {
            "classification": report["classification"],
            "diagnosis_class": report["diagnosis_class"],
            "severity": report["severity"],
            "confidence": report["confidence"],
            "statement": report["summary"],
        },
        "timeline": timeline_rows,
        "errors": errors,
        "hardware_state": hardware_state,
        "cash_state": cash_state,
        "host_state": host_state,
        "analysis": analysis,
        "evidence": evidence,
        "vendor_questions": payload.get("vendor_questions") or [],
        "recommended_actions": payload.get("recommended_actions") or [],
        "traceability": {
            "digest_sha256": explanation["digest_sha256"],
            "provider": explanation["provider"],
            "evidence_count": len(evidence),
            "evidence_bound": BOUNDS["max_evidence"],
            "note": (
                "Every conclusion in this report references file:line evidence "
                "(EV ids). Confidence labels express evidence strength, never "
                "certainty; possible causes are hypotheses."
            ),
        },
    }


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

_SEVERITY_COLORS = {
    "CRITICAL": "#B91C1C",
    "ERROR": "#B91C1C",
    "WARNING": "#B45309",
    "INFO": "#334155",
}


def build_pdf(report: dict) -> bytes:
    """Render the vendor report as a professional PDF (reportlab)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    base = styles["BodyText"]
    base.fontName = "Helvetica"
    base.fontSize = 8.5
    base.leading = 11
    small = ParagraphStyle("small", parent=base, fontSize=7, leading=9, textColor=colors.HexColor("#475569"))
    mono = ParagraphStyle("mono", parent=base, fontName="Courier", fontSize=7, leading=9)
    h1 = ParagraphStyle("h1x", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.white)
    h2 = ParagraphStyle(
        "h2x", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11,
        textColor=colors.HexColor("#1E293B"), spaceBefore=10, spaceAfter=4,
    )

    def esc(text) -> str:
        return sax.escape(str(text)) if text is not None else "—"

    def table(header: list[str], rows: list[list[str]], widths: list[float], mono_cols: tuple[int, ...] = ()) -> Table:
        data = [[Paragraph(f"<b>{esc(h)}</b>", small) for h in header]]
        for row in rows:
            data.append(
                [
                    Paragraph(esc(cell), mono if i in mono_cols else small)
                    for i, cell in enumerate(row)
                ]
            )
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        return t

    story: list = []

    # ---- header band -------------------------------------------------------
    machine = report["machine"] or {}
    txn = report["transaction"] or {}
    problem = report["problem"] or {}
    header_inner = Table(
        [
            [Paragraph("CDM LOG ANALYZER", h1)],
            [
                Paragraph(
                    f"Vendor Escalation Report — Incident {esc(txn.get('transaction_id'))} · "
                    f"generated {esc(report['generated_at'])}",
                    ParagraphStyle("sub", parent=small, textColor=colors.HexColor("#C7D2FE")),
                )
            ],
        ],
        colWidths=[170 * mm],
    )
    header_inner.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#312E81")),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (0, 0), 8),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
            ]
        )
    )
    story.append(header_inner)
    story.append(Spacer(1, 6 * mm))

    # ---- incident summary ----------------------------------------------------
    story.append(Paragraph("Incident Summary", h2))
    meta = Table(
        [
            ["Machine", machine.get("serial_number") or "—", "Model", f"{machine.get('model') or '—'} {machine.get('model_name') or ''}".strip()],
            ["Location", machine.get("location") or "—", "Vendor", machine.get("vendor") or "—"],
            ["Transaction", txn.get("transaction_id") or "—", "Status", txn.get("status") or "—"],
            ["Started", _fmt_dt(_parse_iso(txn.get("start_time"))), "Ended", _fmt_dt(_parse_iso(txn.get("end_time")))],
            [
                "Amount",
                f"{txn.get('amount'):.2f} {txn.get('currency') or ''}".strip() if txn.get("amount") is not None else "—",
                "Classification",
                f"{problem.get('classification')} · {problem.get('diagnosis_class')} · severity {problem.get('severity')} · confidence {problem.get('confidence')}",
            ],
        ],
        colWidths=[26 * mm, 59 * mm, 26 * mm, 59 * mm],
    )
    meta.setStyle(_kv_style())
    story.append(meta)
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(esc(problem.get("statement")), base))
    ai = (report["analysis"] or {}).get("ai_explanation") or {}
    if ai.get("technical_summary"):
        story.append(Spacer(1, 1 * mm))
        story.append(Paragraph(esc(ai["technical_summary"]), small))
    story.append(Spacer(1, 3 * mm))

    # ---- timeline ---------------------------------------------------------------
    story.append(Paragraph("Timeline", h2))
    rows = [
        [
            t["timestamp"] or "—",
            t["event"] + (" (not confirmed)" if t["not_confirmed"] else ""),
            t["device"] or "—",
            t["source"] or "—",
            t["evidence_ref"] or "—",
        ]
        for t in report["timeline"]
    ]
    story.append(table(["Time", "Event", "Device", "Source", "Evidence"], rows, [24 * mm, 52 * mm, 22 * mm, 16 * mm, 56 * mm], mono_cols=(4,)))

    # ---- cash trace ---------------------------------------------------------------
    story.append(Paragraph("Cash Trace", h2))
    cash = report["cash_state"] or {}
    story.append(
        Paragraph(
            f"Final cash state: <b>{esc(cash.get('final_cash_state'))}</b>",
            base,
        )
    )
    rows = [
        [
            m.get("note_id") or "—",
            f"{m.get('from_state')} → {m.get('to_state')}",
            m.get("timestamp") or "—",
            m.get("device") or "—",
            f"{(m.get('confidence') or 0) * 100:.0f}%",
            m.get("evidence_ref") or "—",
        ]
        for m in cash.get("movements") or []
    ] or [["—", "no cash movements recorded", "—", "—", "—", "—"]]
    story.append(table(["Note", "Transition", "Time", "Device", "Conf.", "Evidence"], rows, [22 * mm, 40 * mm, 28 * mm, 24 * mm, 12 * mm, 44 * mm], mono_cols=(5,)))

    # ---- hardware events --------------------------------------------------------------
    story.append(Paragraph("Hardware Events", h2))
    hw = report["hardware_state"] or {}
    hw_rows: list[list[str]] = []
    for s in hw.get("sensors") or []:
        hw_rows.append(["sensor", s.get("sensor") or "—", f"{s.get('transition')}" + (f" (expected {s.get('expected')})" if s.get("expected") else ""), s.get("evidence_ref") or "—"])
    for m in hw.get("motors") or []:
        hw_rows.append(["motor", m.get("motor") or "—", f"duration {m.get('duration_ms') or '—'} ms" + (" TIMEOUT" if m.get("timed_out") else ""), m.get("evidence_ref") or "—"])
    for g in hw.get("gates") or []:
        hw_rows.append([g.get("kind") or "gate", g.get("name") or "—", f"expected {g.get('expected') or '—'} / actual {g.get('actual') or '—'}" + (" MISMATCH" if g.get("state_mismatch") else "") + (" TIMEOUT" if g.get("timed_out") else ""), g.get("evidence_ref") or "—"])
    for t in hw.get("transports") or []:
        hw_rows.append(["transport", t.get("name") or "—", f"outcome {t.get('outcome') or '—'}", t.get("evidence_ref") or "—"])
    if not hw_rows:
        hw_rows = [["—", "—", "no hardware events recorded", "—"]]
    story.append(table(["Kind", "Subject", "Observation", "Evidence"], hw_rows, [18 * mm, 30 * mm, 88 * mm, 34 * mm], mono_cols=(3,)))
    for f in hw.get("faults") or []:
        story.append(Spacer(1, 1 * mm))
        story.append(
            Paragraph(
                f"<b>{esc(f.get('classification'))}</b> — {esc(f.get('statement'))}",
                ParagraphStyle("fault", parent=small, textColor=colors.HexColor("#B91C1C")),
            )
        )

    # ---- analysis -----------------------------------------------------------------------
    story.append(Paragraph("Analysis", h2))
    story.append(Paragraph(esc((report["analysis"] or {}).get("engine_summary")), base))
    rc = ai.get("root_cause") or {}
    if rc:
        story.append(
            Paragraph(
                f"<b>Root cause ({esc(rc.get('label'))})</b>: {esc(rc.get('statement'))} "
                f"<i>[evidence: {esc(', '.join(rc.get('evidence_ids') or []) or '—')}]</i>",
                base,
            )
        )
    for cause in ai.get("possible_causes") or []:
        story.append(
            Paragraph(
                f"Possible cause ({esc(cause.get('label'))}): {esc(cause.get('statement'))}"
                + (f" <i>[{esc(', '.join(cause.get('evidence_ids') or []))}]</i>" if cause.get("evidence_ids") else ""),
                small,
            )
        )
    story.append(
        Paragraph(
            f"Generator: {esc(ai.get('provider'))} · digest {esc((report['traceability'] or {}).get('digest_sha256'))}",
            small,
        )
    )

    # ---- recommendations ------------------------------------------------------------
    story.append(Paragraph("Recommendations", h2))
    for action in report.get("recommended_actions") or ["—"]:
        story.append(Paragraph(f"• {esc(action)}", base))

    # ---- vendor questions ---------------------------------------------------------
    story.append(Paragraph("Vendor Questions", h2))
    for q in report.get("vendor_questions") or ["—"]:
        story.append(Paragraph(f"? {esc(q)}", base))

    # ---- evidence appendix -----------------------------------------------------------
    story.append(Paragraph("Evidence", h2))
    rows = [
        [e["id"], f"{e['file']}:{e['line_number']}", e["origin"], e["raw_excerpt"] or "—"]
        for e in report["evidence"]
    ]
    story.append(table(["ID", "File:Line", "Origin", "Raw excerpt (verbatim)"], rows, [16 * mm, 48 * mm, 22 * mm, 84 * mm], mono_cols=(1, 3)))

    story.append(Spacer(1, 4 * mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#CBD5E1")))
    story.append(
        Paragraph(
            esc((report["traceability"] or {}).get("note"))
            + " Generated by "
            + esc(report["system"])
            + ".",
            small,
        )
    )

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(
            15 * mm,
            8 * mm,
            f"CDM Log Analyzer — vendor escalation report — transaction {txn.get('transaction_id')}",
        )
        canvas.drawRightString(195 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=f"Vendor Escalation Report {txn.get('transaction_id')}",
    )
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def _parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def _kv_style():
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle

    return TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E2E8F0")),
            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E2E8F0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
    )


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

_SHEETS = [
    "Transaction Summary",
    "Events",
    "Cash Trace",
    "Hardware Events",
    "Errors",
    "Analysis",
]


def build_xlsx(report: dict) -> bytes:
    """Render the vendor report as a structured Excel workbook (openpyxl)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    wrap = Alignment(vertical="top", wrap_text=True)

    def sheet(name: str, title_row: list[str], rows: list[list])  :
        ws = wb.create_sheet(name[:31])
        ws.append(title_row)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = wrap
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(title_row))}{max(ws.max_row, 1)}"
        for idx, width in enumerate(_widths(title_row), start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        return ws

    def _widths(title_row: list[str]) -> list[float]:
        return [min(max(len(h) + 4, 12), 60) for h in title_row]

    machine = report["machine"] or {}
    txn = report["transaction"] or {}
    problem = report["problem"] or {}
    ai = (report["analysis"] or {}).get("ai_explanation") or {}
    trace = report["traceability"] or {}

    # 1. Transaction Summary ----------------------------------------------------
    ws = wb.active
    ws.title = _SHEETS[0]
    summary_rows: list[tuple[str, object]] = [
        ("Report kind", report["report_kind"]),
        ("Generated at", report["generated_at"]),
        ("System", report["system"]),
        ("Machine serial", machine.get("serial_number")),
        ("Machine name", machine.get("name")),
        ("Location", machine.get("location")),
        ("Model", machine.get("model")),
        ("Model name", machine.get("model_name")),
        ("Vendor", machine.get("vendor")),
        ("Transaction ID", txn.get("transaction_id")),
        ("Started", txn.get("start_time")),
        ("Ended", txn.get("end_time")),
        ("Amount", txn.get("amount")),
        ("Currency", txn.get("currency")),
        ("Status", txn.get("status")),
        ("Problem classification", problem.get("classification")),
        ("Diagnosis class", problem.get("diagnosis_class")),
        ("Severity", problem.get("severity")),
        ("Confidence", problem.get("confidence")),
        ("Problem statement", problem.get("statement")),
        ("AI technical summary", ai.get("technical_summary")),
        ("AI root cause", (ai.get("root_cause") or {}).get("statement")),
        ("Root cause label", (ai.get("root_cause") or {}).get("label")),
        ("Explanation provider", ai.get("provider")),
        ("Digest SHA-256", trace.get("digest_sha256")),
        ("Evidence references", trace.get("evidence_count")),
        ("Traceability note", trace.get("note")),
    ]
    ws.append(["Field", "Value"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for k, v in summary_rows:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    # 2. Events ---------------------------------------------------------------------
    sheet(
        _SHEETS[1],
        ["Timestamp", "Event", "Stage", "Device", "Severity", "Source", "Evidence", "Raw excerpt"],
        [
            [
                t.get("timestamp"),
                t.get("event") + (" (not confirmed)" if t.get("not_confirmed") else ""),
                t.get("stage"),
                t.get("device"),
                t.get("severity"),
                t.get("source"),
                t.get("evidence_ref"),
                t.get("raw_excerpt"),
            ]
            for t in report["timeline"]
        ],
    )

    # 3. Cash Trace -----------------------------------------------------------------
    sheet(
        _SHEETS[2],
        ["Note", "From", "To", "Timestamp", "Device", "Confidence", "Evidence"],
        [
            [
                m.get("note_id"),
                m.get("from_state"),
                m.get("to_state"),
                m.get("timestamp"),
                m.get("device"),
                m.get("confidence"),
                m.get("evidence_ref"),
            ]
            for m in (report["cash_state"] or {}).get("movements") or []
        ],
    )

    # 4. Hardware Events ----------------------------------------------------------------
    hw_rows: list[list] = []
    for s in (report["hardware_state"] or {}).get("sensors") or []:
        hw_rows.append(["sensor", s.get("sensor"), s.get("transition"), s.get("expected"), s.get("actual"), s.get("abnormal_duration_ms"), s.get("evidence_ref")])
    for m in (report["hardware_state"] or {}).get("motors") or []:
        hw_rows.append(["motor", m.get("motor"), f"duration {m.get('duration_ms')} ms", None, "TIMEOUT" if m.get("timed_out") else None, m.get("timeout_ms"), m.get("evidence_ref")])
    for g in (report["hardware_state"] or {}).get("gates") or []:
        hw_rows.append([g.get("kind"), g.get("name"), f"expected {g.get('expected')} / actual {g.get('actual')}", g.get("expected"), "MISMATCH" if g.get("state_mismatch") else ("TIMEOUT" if g.get("timed_out") else None), None, g.get("evidence_ref")])
    for t in (report["hardware_state"] or {}).get("transports") or []:
        hw_rows.append(["transport", t.get("name"), f"outcome {t.get('outcome')}", None, None, t.get("timeout_ms"), t.get("evidence_ref")])
    sheet(
        _SHEETS[3],
        ["Kind", "Subject", "Observation", "Expected", "Anomaly", "Timeout ms", "Evidence"],
        hw_rows,
    )

    # 5. Errors ------------------------------------------------------------------------------
    sheet(
        _SHEETS[4],
        ["Timestamp", "Code", "Description", "Event", "Device", "Evidence", "Raw excerpt"],
        [
            [
                e.get("timestamp"),
                e.get("error_code"),
                e.get("description"),
                e.get("event"),
                e.get("device"),
                e.get("evidence_ref"),
                e.get("raw_excerpt"),
            ]
            for e in report["errors"]
        ],
    )

    # 6. Analysis ------------------------------------------------------------------------------
    ws = wb.create_sheet(_SHEETS[5])
    ws.append(["Section", "Rule / Item", "Label / Severity", "Confidence", "Statement", "Evidence"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap
    for f in (report["analysis"] or {}).get("findings") or []:
        ws.append(
            [
                "Rule finding",
                f.get("rule_id"),
                f"{f.get('severity')} / {f.get('diagnosis_class')}",
                f.get("confidence"),
                f.get("summary"),
                "; ".join(f.get("evidence_refs") or []),
            ]
        )
        if f.get("interpretation"):
            ws.append(["Interpretation", f.get("rule_id"), None, None, f.get("interpretation"), None])
        for cause in f.get("possible_causes") or []:
            ws.append(["Possible cause", f.get("rule_id"), "POSSIBLE", None, cause if isinstance(cause, str) else str(cause), None])
    rc = ai.get("root_cause") or {}
    ws.append(["AI root cause", ai.get("provider"), rc.get("label"), (ai.get("confidence") or {}).get("label"), rc.get("statement"), "; ".join(rc.get("evidence_ids") or [])])
    for cause in ai.get("possible_causes") or []:
        ws.append(["AI possible cause", ai.get("provider"), cause.get("label"), None, cause.get("statement"), "; ".join(cause.get("evidence_ids") or [])])
    for q in report.get("vendor_questions") or []:
        ws.append(["Vendor question", None, None, None, q, None])
    for note in ai.get("safety_notes") or []:
        ws.append(["Safety note", None, None, None, note, None])
    ws.freeze_panes = "A2"
    for idx, width in enumerate([18, 26, 22, 14, 90, 40], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
